# -*- coding: utf-8 -*-
r"""Paper 1 -- Response to Reviewers Markdown -> Word (.docx)"""
import os, re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = r'D:/WorkBuddy/2026-10-06-10-35-00/outputs/paper1'
MD = os.path.join(BASE, 'Paper1_Response_to_Reviewers_中文.md')
OUT = os.path.join(BASE, 'Paper1_Response_to_Reviewers_答复函.docx')

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin = sec.bottom_margin = Cm(2.54)
sec.left_margin = sec.right_margin = Cm(2.4)

style = doc.styles['Normal']
style.font.name = 'Times New Roman'; style.font.size = Pt(11.5)
style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')


def set_cn(run, east='宋体', west='Times New Roman'):
    run.font.name = west
    run._element.rPr.rFonts.set(qn('w:eastAsia'), east)


def add_para(text='', size=11.5, bold=False, align=None, east='宋体', space_after=5,
             line=1.4, indent_first=False, color=None):
    text = text.replace('`', '')
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    if line:
        pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE; pf.line_spacing = line
    if align is not None:
        p.alignment = align
    if indent_first:
        pf.first_line_indent = Pt(size * 2)
    parts = re.split(r'(\*\*.+?\*\*)', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            r = p.add_run(part[2:-2]); r.bold = True
        else:
            r = p.add_run(part); r.bold = bold
        r.font.size = Pt(size); set_cn(r, east)
        if color:
            r.font.color.rgb = RGBColor(*color)
    return p


def add_heading(text, level):
    if level == 1:
        add_para(text, size=15, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, east='黑体', space_after=10)
    elif level == 2:
        add_para(text, size=13, bold=True, east='黑体', space_after=7)
    else:
        add_para(text, size=11.5, bold=True, east='黑体', space_after=5)


def set_cell_border(cell, top=None, bottom=None):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    for tag, sz in (('top', top), ('bottom', bottom)):
        if sz:
            el = OxmlElement('w:' + tag)
            el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), str(sz))
            el.set(qn('w:color'), '000000'); el.set(qn('w:space'), '0')
            borders.append(el)
    tcPr.append(borders)


def add_table(rows, font_size=9):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            c = t.cell(i, j); c.text = ''
            p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(1)
            r = p.add_run(str(cell)); r.font.size = Pt(font_size); set_cn(r)
            if i == 0:
                r.bold = True; set_cell_border(c, top=12)
            if i == len(rows) - 1:
                set_cell_border(c, bottom=12)
            if i == 1:
                set_cell_border(c, top=6)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)


lines = open(MD, encoding='utf-8').read().split('\n')
i = 0
code_buf = None
while i < len(lines):
    ln = lines[i].rstrip()
    # 代码块
    if ln.strip().startswith('```'):
        if code_buf is None:
            code_buf = []
        else:
            for cl in code_buf:
                p = add_para(cl, size=9.5, east='Consolas',
                             space_after=0, line=1.15)
                p.paragraph_format.left_indent = Cm(0.8)
            code_buf = None
        i += 1; continue
    if code_buf is not None:
        code_buf.append(ln); i += 1; continue

    if ln.startswith('# '):
        add_heading(ln[2:].strip(), 1); i += 1; continue
    if ln.startswith('### '):
        add_heading(ln[4:].strip(), 3); i += 1; continue
    if ln.startswith('## '):
        add_heading(ln[3:].strip(), 2); i += 1; continue
    if ln.startswith('> '):
        add_para(ln[2:].strip(), size=10.5, color=(0x59, 0x59, 0x59), space_after=3, line=1.3,
                 indent_first=False)
        i += 1; continue
    if ln.strip() in ('---', ''):
        i += 1; continue
    if ln.startswith('|'):
        rows = []
        while i < len(lines) and lines[i].strip().startswith('|'):
            cells = [c.strip() for c in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r':?-{2,}:?', c) for c in cells):
                rows.append(cells)
            i += 1
        if rows:
            add_table(rows)
        continue
    if ln.startswith('- ') or ln.startswith('* '):
        p = add_para('• ' + ln[2:], size=11, space_after=3, line=1.35)
        p.paragraph_format.left_indent = Cm(0.6)
        i += 1; continue
    m = re.match(r'^(\d+)\. (.+)', ln)
    if m:
        p = add_para(ln, size=11, space_after=3, line=1.35)
        p.paragraph_format.left_indent = Cm(0.7)
        p.paragraph_format.first_line_indent = Cm(-0.7)
        i += 1; continue
    add_para(ln, size=11.5, space_after=5, line=1.4)
    i += 1

doc.save(OUT)
print('saved:', OUT)
