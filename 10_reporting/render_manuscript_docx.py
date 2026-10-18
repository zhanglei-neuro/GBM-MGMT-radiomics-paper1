# -*- coding: utf-8 -*-
r"""Paper 1 - Markdown → 学术排版 Word (.docx)
排版: A4 / 正文小四宋体 1.5倍行距 / 西文 Times New Roman / 三线表 / 300dpi 插图
"""
import os, re, csv
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = r'D:/WorkBuddy/2026-10-06-10-35-00/outputs/paper1'
MD   = os.path.join(BASE, 'Paper1_稿件_中文完整版.md')
FIG  = os.path.join(BASE, 'figures')
ANA  = os.path.join(BASE, 'analysis')
OUT  = os.path.join(BASE, 'Paper1_MGMT_multivendor_radiomics_中文稿件.docx')

doc = Document()
# 页面
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin = sec.bottom_margin = Cm(2.54)
sec.left_margin = sec.right_margin = Cm(2.6)

style = doc.styles['Normal']
style.font.name = 'Times New Roman'; style.font.size = Pt(12)
style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

def set_cn(run, east='宋体', west='Times New Roman'):
    run.font.name = west
    run._element.rPr.rFonts.set(qn('w:eastAsia'), east)

def add_para(text='', size=12, bold=False, align=None, east='宋体', space_after=6,
             line=1.5, indent_first=False, color=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    if line: pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE; pf.line_spacing = line
    if align is not None: p.alignment = align
    if indent_first: pf.first_line_indent = Pt(size * 2)
    # **bold** inline
    parts = re.split(r'(\*\*.+?\*\*)', text)
    for part in parts:
        if not part: continue
        if part.startswith('**') and part.endswith('**'):
            r = p.add_run(part[2:-2]); r.bold = True
        else:
            r = p.add_run(part); r.bold = bold
        r.font.size = Pt(size); set_cn(r, east)
        if color: r.font.color.rgb = RGBColor(*color)
    return p

def add_heading(text, level):
    if level == 1:
        add_para(text, size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, east='黑体', space_after=10)
    elif level == 2:
        add_para(text, size=14, bold=True, east='黑体', space_after=8)
    else:
        add_para(text, size=12, bold=True, east='黑体', space_after=6)

def set_cell_border(cell, top=None, bottom=None):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    for tag, sz in (('top', top), ('bottom', bottom)):
        if sz:
            el = OxmlElement('w:' + tag)
            el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), str(sz))
            el.set(qn('w:color'), '000000'); el.set(qn('w:space'), '0')
            borders.append(el)
    tcPr.append(borders)

def add_table(rows, caption=None, font_size=10.5):
    if caption:
        add_para(caption, size=10.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4, line=1.2)
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            c = t.cell(i, j); c.text = ''
            p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(str(cell).replace('**', '')); r.font.size = Pt(font_size); set_cn(r)
            if i == 0: r.bold = True
            if i == 0: set_cell_border(c, top=12)
            elif i == len(rows) - 1: set_cell_border(c, bottom=12)
            elif i == 1: set_cell_border(c, top=6)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_figure(path, caption, width_cm=15):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(path, width=Cm(width_cm))
    add_para(caption, size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12, line=1.2)

# ---------------- 解析 md ----------------
lines = open(MD, encoding='utf-8').read().split('\n')
i = 0
title_done = False
in_ref = False
while i < len(lines):
    ln = lines[i].rstrip()
    if ln.startswith('# ') and not title_done:
        add_heading(ln[2:].strip(), 1); title_done = True; i += 1; continue
    if ln.startswith('### '):
        add_heading(ln[4:].strip(), 3); i += 1; continue
    if ln.startswith('## '):
        txt = ln[3:].strip()
        in_ref = txt.startswith('参考文献')
        add_heading(txt, 2); i += 1; continue
    if ln.startswith('> '):
        add_para(ln[2:].strip(), size=10.5, color=(0x59, 0x59, 0x59), space_after=3, line=1.3)
        i += 1; continue
    if ln.strip() == '---' or ln.strip() == '':
        i += 1; continue
    if ln.startswith('|'):
        # 表格块
        rows = []
        while i < len(lines) and lines[i].strip().startswith('|'):
            cells = [c.strip() for c in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r':?-{2,}:?', c) for c in cells):
                rows.append(cells)
            i += 1
        if rows: add_table(rows)
        continue
    m = re.match(r'^(\d+)\. (.+)', ln)
    if m and in_ref:
        p = add_para(ln, size=10.5, space_after=3, line=1.3)
        p.paragraph_format.left_indent = Cm(0.75)
        p.paragraph_format.first_line_indent = Cm(-0.75)
        i += 1; continue
    if ln.startswith('- '):
        add_para('• ' + ln[2:], size=11, space_after=3, line=1.4)
        i += 1; continue
    # 英文标题行
    if ln.startswith('**Multivendor'):
        add_para(ln.strip('*'), size=13, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)
        i += 1; continue
    # 正文
    add_para(ln, size=12, indent_first=True, space_after=6)
    i += 1

# ---------------- 图表页 ----------------
doc.add_page_break()
add_heading('图表', 2)

# 表 1
rows = list(csv.reader(open(os.path.join(ANA, 'table1_main.csv'), encoding='utf-8-sig')))
add_table(rows, caption='表 1  队列基线特征（总体与按 MRI 厂商分层）', font_size=9)
add_para('注：连续变量为中位数 [四分位距]，分类变量为 n (%)。p 值为 Siemens 与 GE 两亚组的组间检验'
         '（连续变量用 Mann–Whitney U 检验，分类变量用 χ² 或 Fisher 精确检验）。本表纳入的全部临床变量'
         '（年龄、性别、KPS、切除程度、肿瘤部位、肿瘤侧别）在 89 例中均完整记录，不存在缺失，故各变量分母一致；'
         '完整版见补充表 S5。全队列 89 例均为 IDH 野生型 GBM（WHO 2021 定义实体），故 IDH 不设组间比较。'
         '肿瘤侧别左 41、右 39、双侧/中线 9。', size=8.5, line=1.15, space_after=10)

# 表 2
rows = list(csv.reader(open(os.path.join(ANA, 'table2_performance.csv'), encoding='utf-8-sig')))
add_table(rows, caption='表 2  重复嵌套交叉验证（5 折 × 5 次重复）下各模型性能')
add_para('注：影像组学与临床基线模型的 AUC 为 25 个外折的重复均值 ± 标准差；95%CI 基于合并 OOF 预测的 2000 次 bootstrap。'
         '「临床+影像组学（联合）」行的 AUC/95%CI/Brier 与「影像组学—逻辑回归」完全相同，系因年龄/性别在 25 个外折中'
         '从未进入最终特征集（见正文 3.3 节机制核验），两模型逐例预测概率 bit 级一致（max|Δp|=0），并非数据错误。',
         size=8.5, line=1.15, space_after=10)

# 表 3
rows = list(csv.reader(open(os.path.join(ANA, 'table3_main.csv'), encoding='utf-8-sig')))
add_table(rows, caption='表 3  体积生物标志物与 MGMT 状态的关联及单变量 AUC', font_size=9)
add_para('注：体积取自各亚区掩膜在其原生网格上的逐体素计数（n=89，无缺失）。AUC 方向统一为'
         '「指标升高 → 预测 MGMT 启动子甲基化（MGMT+）」，即 AUC>0.5 表示甲基化组该指标更高，'
         '与「Higher in」列一致。q 值为 7 个预设体积标志物族内的 Benjamini–Hochberg FDR 校正值；'
         '全部 7 项校正后均不显著（最小 q=0.060）。', size=8.5, line=1.15, space_after=10)

figs = [
    ('fig1_batch_effect.png',
     '图 1  4148 维组学特征的 PCA 与 t-SNE 嵌入。上排按扫描仪厂商着色：Siemens 与 GE 呈明显分离；下排按 MGMT 状态着色：两组未见分离趋势。'
     '定量批次指标与之一致：平均轮廓系数（ASW）厂商 0.280（置换 p=0.005）vs MGMT 0.012（p=0.050）；PC1–10 组间变异占比厂商 37.0% vs 标签 2.3%。'
     '注：t-SNE 仅为定性展示，结论以 ASW 与 PC 方差占比为准。'),
    ('fig2_roc.png',
     '图 2  重复分层 5 折嵌套交叉验证（5 次重复）下各模型的 ROC 曲线。全部影像组学模型与临床基线的 AUC 95%CI 均覆盖 0.5；'
     '在合并 OOF 预测上，影像组学逻辑回归（0.602）与临床基线（0.596）无统计学差异（DeLong 检验 p=0.795）。'),
    ('fig3_calibration_dca.png',
     '图 3  模型校准曲线（左）与决策曲线分析（右）。影像组学逻辑回归 OOF 预测的校准斜率 0.570、校准截距 −0.001（n=89），'
     '即无系统性高/低估但预测概率过度分散；模型在全部阈值范围无临床净获益。'),
    ('fig4_vendor_robustness.png',
     '图 4  跨厂商鲁棒性。(a) 留一厂商交叉验证（LOVO，XGBoost）：训练 GE→测试 Siemens 0.450，训练 Siemens→测试 GE 0.524；'
     '(b) 厂商内部嵌套 CV：Siemens logistic 0.579±0.061 / XGBoost 0.565±0.082，GE logistic 0.328±0.094 / XGBoost 0.220±0.055（性能低于随机且方向不稳定）；'
     '(c) ComBat 与厂商内 Z 标准化（VWS）调和后判别力仍未恢复。'),
    ('fig5_feature_vendor.png',
     '图 5  特征层面的厂商效应。(a) 4148 特征的 Kruskal–Wallis p 值分布（BH-FDR q<0.05 者 3382 个，占 81.5%）；'
     '(b) 厂商效应量最大的 15 个特征（全特征中位秩双列效应量 0.767）。'
     '经采集参数去混杂后（厂商 + 序列类型 + 层厚 + TR/TE/TI），FDR 显著特征降至 259/4148（6.2%）。'),
    ('fig6_shap.png',
     '图 6  SHAP 特征重要性（XGBoost）。Exploratory analysis, fitted on the full dataset, not cross-validated。'
     '仅作假设生成的探索性展示，不作为主要结论依据。'),
    ('fig7_roi_auc.png',
     '图 7  ROI×序列分层的探索性 AUC。(a) 热图；(b) 四个主组合对比。'
     'TumorCore–T1CE（0.634±0.028）、FLAIRChanges–FLAIR（0.614±0.014）与 Necrosis–T1CE（0.576±0.021）呈微弱高于随机的趋势，'
     'Edema–FLAIR 无信号（0.482）；全部信号在跨厂商迁移中消失。鉴于 Raidionics 坏死分割与手动金标准一致性有限（ICC 0.50），'
     '凡以坏死区特征为主导的读数均只具探索性证据权重。'),
    ('fig8_volume.png',
     '图 8  体积生物标志物。(a) 总体肿瘤/水肿比（TT/ED）按 MGMT 状态的分布；(b) 各体积标志物单变量 AUC 与 Zahirovic 2025 外部基准（AUC 0.687）对标——'
     '本队列 TT/ED 比 AUC 0.622（95%CI 0.503–0.732），方向与基准一致但更弱；(c) TT/ED 比 AUC 的厂商依赖性：Siemens 0.685（95%CI 0.538–0.818，'
     '与 Zahirovic 的 0.687 几乎相同）vs GE 0.530。7 项标志物经族内 BH-FDR 校正后均不显著（最小 q=0.060）。'),
]
for f, cap in figs:
    add_figure(os.path.join(FIG, f), cap)

doc.save(OUT)
print('saved:', OUT)
